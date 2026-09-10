import os
from pathlib import Path

from fastapi import FastAPI, HTTPException
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel

app = FastAPI(title="CX Training Service")


class ParseRequest(BaseModel):
    file_path: str
    file_type: str


class ChunkRequest(BaseModel):
    text: str
    source_meta: dict


def parse_pdf(filepath: str) -> str:
    import fitz
    doc = fitz.open(filepath)
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text("text").strip()
        if text:
            pages.append(f"[Page {i + 1}]\n{text}")
    return "\n\n".join(pages)


def parse_docx(filepath: str) -> str:
    from docx import Document
    doc = Document(filepath)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def parse_doc(filepath: str) -> str:
    try:
        return parse_docx(filepath)
    except Exception:
        with open(filepath, "rb") as f:
            return f.read().decode("utf-8", errors="ignore")


@app.post("/parse")
async def parse_document(payload: ParseRequest):
    if not os.path.exists(payload.file_path):
        raise HTTPException(404, "File not found")
    ft = payload.file_type.lower()
    if ft == "pdf":
        text = parse_pdf(payload.file_path)
    elif ft == "docx":
        text = parse_docx(payload.file_path)
    elif ft in ("doc", "txt"):
        text = parse_doc(payload.file_path) if ft == "doc" else Path(payload.file_path).read_text(encoding="utf-8", errors="ignore")
    else:
        raise HTTPException(400, f"Unsupported type: {ft}")
    return {"text": text, "char_count": len(text)}


@app.post("/chunk")
async def chunk_text(payload: ChunkRequest):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=512,
        chunk_overlap=64,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_text(payload.text)
    meta = payload.source_meta
    return {
        "chunks": [
            {
                "chunk_id": f"{meta['source_id']}_{i}",
                "client_id": meta["client_id"],
                "source_id": meta["source_id"],
                "source_type": meta["type"],
                "source_name": meta["name"],
                "source_url": meta.get("url", meta.get("filename", "")),
                "text": chunk,
                "chunk_index": i,
            }
            for i, chunk in enumerate(chunks)
            if chunk.strip()
        ]
    }


@app.get("/health")
async def health():
    return {"status": "ok", "service": "training"}
